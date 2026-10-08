"""Lectura de la base de datos en Excel (.xlsx / .xlsm). La fila 1 son los encabezados."""
from __future__ import annotations

import datetime as dt

from openpyxl import load_workbook


def list_sheets(path: str) -> list[str]:
    wb = load_workbook(path, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def _decimals(number_format: str) -> int | None:
    """Decimales fijos de un formato de celda como '0.00' o '$#,##0.00'; None si no aplica."""
    fmt = (number_format or "").split(";")[0]
    if fmt == "General" or "." not in fmt:
        return None
    after = fmt.split(".", 1)[1]
    n = len(after) - len(after.lstrip("0"))
    return n or None


def _fmt(v, number_format: str = "") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, (int, float)) and (dec := _decimals(number_format)) is not None:
        return f"{v:.{dec}f}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y" if v.time() == dt.time(0) else "%d/%m/%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def read_sheet(path: str, sheet: str) -> tuple[list[str], list[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        it = wb[sheet].iter_rows()
        first = next(it, None)
        if first is None:
            return [], []
        headers, seen = [], {}
        for i, h in enumerate((c.value for c in first), start=1):
            name = _fmt(h) or f"Columna{i}"
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 1
            headers.append(name)
        rows = []
        for cells in it:
            values = [_fmt(c.value, getattr(c, "number_format", "")) for c in cells]
            if not any(values):
                continue
            values += [""] * (len(headers) - len(values))
            rows.append(dict(zip(headers, values)))
        return headers, rows
    finally:
        wb.close()
